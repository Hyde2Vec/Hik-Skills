#!/usr/bin/env python3
"""查看工作表、汇总数值、比较记录；业务含义由调用方判断。"""
import argparse
import json
import math
import sys
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
    get_column_letter = None

BLANK = {"", "-", "—", "–", "－"}
SAMPLE_ROWS = 6
SINGLE_SAMPLE_ROWS = 12
SAMPLE_CELLS = 80
OVERVIEW_COLS = 40


def fail(message):
    raise ValueError(message)


def json_safe(value):
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time.min else value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return str(value)


def dump(data, output=None):
    text = json.dumps(data, ensure_ascii=False, indent=2, default=json_safe) + "\n"
    if output:
        path = Path(output).expanduser()
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
    sys.stdout.write(text)


def load_json(spec):
    if spec == "-":
        raw, label = sys.stdin.read(), "标准输入"
    else:
        path = Path(spec).expanduser()
        if path.is_file():
            try:
                raw, label = path.read_text(encoding="utf-8-sig"), str(path)
            except UnicodeDecodeError as exc:
                fail(f"{path}不是UTF-8文本：{exc}")
        else:
            raw, label = spec, "参数"
    try:
        return json.loads(raw), label
    except json.JSONDecodeError as exc:
        if spec != "-" and not str(spec).lstrip().startswith(("[", "{", '"')):
            fail(f"找不到文件：{spec}")
        fail(f"{label}不是合法JSON：{exc.msg}（第{exc.lineno}行第{exc.colno}列）")


def load_records(spec):
    data, label = load_json(spec)
    if not isinstance(data, list):
        fail(f"{label}必须是JSON数组")
    records = []
    for index, item in enumerate(data, 1):
        if not isinstance(item, dict):
            fail(f"{label}第{index}条不是对象")
        records.append(item)
    return records, label


def require_fields(record, fields, loc):
    missing = [name for name in fields if name not in record]
    if missing:
        fail(f"{loc}缺少字段：{'、'.join(missing)}")


def is_number(value):
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def finite_decimal(value, loc):
    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, int):
        number = Decimal(value)
    elif isinstance(value, float):
        if not math.isfinite(value):
            fail(f"{loc}：不是有限数")
        number = Decimal(str(value))
    else:
        fail(f"{loc}：不是数值")
    if not number.is_finite():
        fail(f"{loc}：不是有限数")
    return number


def parse_number(value, loc):
    if value is None:
        return None
    if isinstance(value, bool):
        fail(f"{loc}：布尔值不是数值")
    if is_number(value):
        return finite_decimal(value, loc)
    if isinstance(value, str):
        text = value.strip()
        if text in BLANK:
            return None
        try:
            number = Decimal(text)
        except InvalidOperation:
            fail(f"{loc}：无法识别为数值：{value}")
        if not number.is_finite():
            fail(f"{loc}：不是有限数：{value}")
        return number
    fail(f"{loc}：数值类型不支持：{type(value).__name__}")


def values_equal(left, right):
    if left is None or right is None:
        return left is None and right is None
    if isinstance(left, bool) or isinstance(right, bool):
        return isinstance(left, bool) and isinstance(right, bool) and left == right
    if is_number(left) and is_number(right):
        return finite_decimal(left, "比较") == finite_decimal(right, "比较")
    return left == right


def canon_key(value, loc):
    if value is None:
        return ("empty", None)
    if isinstance(value, bool):
        return ("bool", value)
    if is_number(value):
        return ("num", finite_decimal(value, loc))
    if isinstance(value, (list, dict)):
        return ("json", json.dumps(value, ensure_ascii=False, sort_keys=True, default=str))
    return ("val", value)


def is_blank_identity(value):
    return value is None or (isinstance(value, str) and not value.strip())


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def merge_index(ws):
    origins, ranges = {}, []
    for area in ws.merged_cells.ranges:
        text = str(area)
        ranges.append(text)
        origin = (area.min_row, area.min_col)
        for row in range(area.min_row, area.max_row + 1):
            for column in range(area.min_col, area.max_col + 1):
                origins[row, column] = (origin, text)
    return origins, ranges


def read_cell(formula_ws, cached_ws, origins, row, column):
    origin, merge = origins.get((row, column), ((row, column), None))
    raw = formula_ws.cell(*origin).value
    cached = cached_ws.cell(*origin).value if cached_ws is not None else None
    if is_formula(cached):
        cached = None
    item = {
        "行": row,
        "列": column,
        "坐标": f"{get_column_letter(column)}{row}",
        "原始": json_safe(raw),
        "缓存": json_safe(cached),
    }
    if is_formula(raw) and cached is None:
        item["无缓存公式"] = True
    if merge:
        item["合并区域"] = merge
        if origin != (row, column):
            item["合并到"] = f"{get_column_letter(origin[1])}{origin[0]}"
    return item, raw, cached


def occupied(raw, cached):
    return raw is not None or cached is not None


def open_workbook(path):
    if openpyxl is None:
        fail("读取 Excel 需要 openpyxl")
    file = Path(path).expanduser()
    if not file.is_file():
        fail(f"找不到 Excel 文件：{path}")
    if file.suffix.lower() not in {".xlsx", ".xlsm"}:
        fail(f"当前只读取 .xlsx/.xlsm：{path}")
    try:
        formulas = openpyxl.load_workbook(file, data_only=False)
    except Exception as exc:
        fail(f"Excel 无法打开：{file}（{exc}）")
    try:
        cached = openpyxl.load_workbook(file, data_only=True)
    except Exception as exc:
        formulas.close()
        fail(f"Excel 缓存值无法打开：{file}（{exc}）")
    return file, formulas, cached


def sheet_by_name(wb, name):
    if name not in wb.sheetnames:
        fail(f"找不到工作表：{name}。现有：{'、'.join(wb.sheetnames) or '无'}")
    ws = wb[name]
    if not hasattr(ws, "merged_cells"):
        fail(f"{name} 不是可读取的工作表")
    return ws


def collect_cells(formula_ws, cached_ws, origins, rows, columns):
    cells = []
    for row in rows:
        for column in columns:
            item, raw, cached = read_cell(formula_ws, cached_ws, origins, row, column)
            if occupied(raw, cached):
                cells.append(item)
    return cells


def preview_rows(max_row, sample_n):
    if max_row <= 0:
        return []
    first = list(range(1, min(max_row, sample_n) + 1))
    last = [row for row in (max_row - 1, max_row) if row > first[-1]]
    return first + last


def overlapping_merges(ws, row_start, row_end, column_start, column_end):
    ranges = []
    for area in ws.merged_cells.ranges:
        if area.max_row < row_start or area.min_row > row_end or area.max_col < column_start or area.min_col > column_end:
            continue
        ranges.append(str(area))
    return ranges


def inspect_sheet(formula_ws, cached_ws, start=None, end=None, col_start=None, col_end=None, full=False, sample_n=SAMPLE_ROWS):
    origins, merges = merge_index(formula_ws)
    max_row = formula_ws.max_row or 0
    max_col = formula_ws.max_column or 0
    info = {
        "名称": formula_ws.title,
        "行数": max_row,
        "列数": max_col,
        "隐藏": formula_ws.sheet_state != "visible",
        "合并区域": merges,
    }
    if full:
        if any(value is not None and value < 1 for value in (start, end, col_start, col_end)):
            fail("行列范围从 1 开始")
        if start > end:
            fail("起始行不能大于结束行")
        if col_start is not None and col_end is not None and col_start > col_end:
            fail("起始列不能大于结束列")
        column_start = 1 if col_start is None else col_start
        column_end = max_col if col_end is None else col_end
        info["范围"] = {"起始行": start, "结束行": end, "起始列": column_start, "结束列": column_end}
        info["合并区域"] = overlapping_merges(formula_ws, start, end, column_start, column_end)
        row_from, row_to = max(1, start), min(end, max_row)
        col_from, col_to = max(1, column_start), min(column_end, max_col)
        if row_from > row_to or col_from > col_to:
            info["单元格"] = []
        else:
            info["单元格"] = collect_cells(
                formula_ws, cached_ws, origins,
                range(row_from, row_to + 1), range(col_from, col_to + 1),
            )
        info["空单元格已省略"] = True
        return info
    columns = min(max_col, OVERVIEW_COLS)
    preview = preview_rows(max_row, sample_n)
    cells = collect_cells(formula_ws, cached_ws, origins, preview, range(1, columns + 1)) if columns else []
    info["样例行"] = preview
    info["样例"] = cells[:SAMPLE_CELLS]
    if max_col > OVERVIEW_COLS:
        info["样例仅前列"] = OVERVIEW_COLS
    if len(cells) > SAMPLE_CELLS:
        info["样例已截断"] = True
    return info


def cmd_inspect(args):
    path, formulas, cached = open_workbook(args.输入)
    cached_map = {ws.title: ws for ws in cached.worksheets}
    range_requested = any(value is not None for value in (args.起始行, args.结束行, args.起始列, args.结束列))
    try:
        if range_requested:
            if not args.工作表 or args.起始行 is None or args.结束行 is None:
                fail("读取原始区域需要同时指定 --工作表、--起始行 和 --结束行")
            ws = sheet_by_name(formulas, args.工作表)
            result = {
                "文件": str(path),
                "工作表": [inspect_sheet(
                    ws, cached_map.get(ws.title),
                    start=args.起始行, end=args.结束行,
                    col_start=args.起始列, col_end=args.结束列, full=True,
                )],
            }
        else:
            sheets = [sheet_by_name(formulas, args.工作表)] if args.工作表 else formulas.worksheets
            sample_n = SINGLE_SAMPLE_ROWS if args.工作表 else SAMPLE_ROWS
            result = {
                "文件": str(path),
                "工作表": [inspect_sheet(ws, cached_map.get(ws.title), sample_n=sample_n) for ws in sheets],
            }
    finally:
        formulas.close()
        cached.close()
    dump(result, args.输出)


def summarize_records(indexed, field, label):
    total, counted, missing = Decimal(0), 0, 0
    for index, record in indexed:
        loc = f"{label}第{index}条字段{field}"
        require_fields(record, [field], loc)
        number = parse_number(record[field], loc)
        if number is None:
            missing += 1
        else:
            total += number
            counted += 1
    return {"有值合计": format(total, "f"), "有值条数": counted, "缺失条数": missing}


def cmd_summarize(args):
    records, label = load_records(args.输入)
    groups, order = {}, []
    indexed = []
    for index, record in enumerate(records, 1):
        loc = f"{label}第{index}条"
        require_fields(record, args.分组, loc)
        key = tuple(canon_key(record[name], f"{loc}字段{name}") for name in args.分组)
        if key not in groups:
            order.append(key)
            groups[key] = []
        item = (index, record)
        indexed.append(item)
        groups[key].append(item)
    overall = {field: summarize_records(indexed, field, label) for field in args.求和}
    grouped = []
    for key in order:
        items = groups[key]
        sample = items[0][1]
        grouped.append({
            "键": {name: json_safe(sample[name]) for name in args.分组},
            "记录数": len(items),
            "字段": {field: summarize_records(items, field, label) for field in args.求和},
        })
    dump({
        "记录数": len(records),
        "分组字段": args.分组,
        "求和字段": args.求和,
        "全体": overall,
        "分组": grouped,
    }, args.输出)


def bucket(records, keys, fields, label):
    groups, order = {}, []
    for index, record in enumerate(records, 1):
        loc = f"{label}第{index}条"
        require_fields(record, list(keys) + [name for name in fields if name not in keys], loc)
        for name in list(keys) + list(fields):
            value = record[name]
            if is_number(value):
                finite_decimal(value, f"{loc}字段{name}")
        key = tuple(canon_key(record[name], f"{loc}字段{name}") for name in keys)
        if key not in groups:
            order.append(key)
            groups[key] = []
        groups[key].append(record)
    return groups, order


def cmd_compare(args):
    current, current_label = load_records(args.本期)
    previous, previous_label = load_records(args.上期)
    now, now_order = bucket(current, args.键, args.字段, current_label)
    old, old_order = bucket(previous, args.键, args.字段, previous_label)
    added, removed, changed, ambiguous = [], [], [], []
    seen = set()
    for key in now_order + [item for item in old_order if item not in now]:
        if key in seen:
            continue
        seen.add(key)
        after, before = now.get(key, []), old.get(key, [])
        sample = (after or before)[0]
        identity = {name: json_safe(sample[name]) for name in args.键}
        blank = any(is_blank_identity(row[name]) for row in after + before for name in args.键)
        if blank or len(after) > 1 or len(before) > 1:
            item = {"键": identity, "本期条数": len(after), "上期条数": len(before), "本期": after, "上期": before}
            if blank:
                item["原因"] = "身份字段为空"
            ambiguous.append(item)
            continue
        if not before:
            added.append({"键": identity, "本期": after, "上期": before})
        elif not after:
            removed.append({"键": identity, "本期": after, "上期": before})
        else:
            fields = [name for name in args.字段 if not values_equal(before[0][name], after[0][name])]
            if fields:
                changed.append({"键": identity, "变化字段": fields, "本期": after, "上期": before})
    dump({
        "本期记录数": len(current),
        "上期记录数": len(previous),
        "键": args.键,
        "字段": args.字段,
        "新增": added,
        "消失": removed,
        "变化": changed,
        "身份歧义": ambiguous,
    }, args.输出)


def build_parser():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="命令", required=True)
    inspect = commands.add_parser("查看表", help="查看工作表结构，或按行列范围读取原始区域")
    inspect.add_argument("--输入", required=True)
    inspect.add_argument("--工作表")
    inspect.add_argument("--起始行", type=int)
    inspect.add_argument("--结束行", type=int)
    inspect.add_argument("--起始列", type=int)
    inspect.add_argument("--结束列", type=int)
    inspect.add_argument("--输出")
    inspect.set_defaults(func=cmd_inspect)
    summarize = commands.add_parser("汇总", help="按指定字段分组求和；缺失单独计数")
    summarize.add_argument("--输入", required=True)
    summarize.add_argument("--分组", nargs="+", default=[])
    summarize.add_argument("--求和", nargs="+", required=True)
    summarize.add_argument("--输出")
    summarize.set_defaults(func=cmd_summarize)
    compare = commands.add_parser("比较", help="按指定身份键和字段比较两期记录")
    compare.add_argument("--本期", required=True)
    compare.add_argument("--上期", required=True)
    compare.add_argument("--键", nargs="+", required=True)
    compare.add_argument("--字段", nargs="+", required=True)
    compare.add_argument("--输出")
    compare.set_defaults(func=cmd_compare)
    return parser


def main():
    parser = build_parser()
    args = parser.parse_args()
    try:
        args.func(args)
    except (ValueError, OSError, UnicodeDecodeError) as exc:
        parser.exit(1, f"未完成：{exc}\n")


if __name__ == "__main__":
    main()
